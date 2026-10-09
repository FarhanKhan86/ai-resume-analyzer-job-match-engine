from pathlib import Path
import fitz
from docx import Document

MAX_FILE_BYTES = 8 * 1024 * 1024

def extract_resume_text(filename: str, content: bytes) -> str:
    """Extract text from supported resume formats without saving uploads to disk."""
    if len(content) > MAX_FILE_BYTES:
        raise ValueError("File is too large. Maximum supported size is 8 MB.")
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        try:
            with fitz.open(stream=content, filetype="pdf") as doc:
                text = "\n".join(page.get_text() for page in doc)
        except Exception as exc:
            raise ValueError("Could not read this PDF. Try exporting it as a text-based PDF.") from exc
    elif suffix == ".docx":
        import io
        try:
            doc = Document(io.BytesIO(content))
            text = "\n".join(p.text for p in doc.paragraphs)
            for table in doc.tables:
                for row in table.rows:
                    text += "\n" + " | ".join(cell.text for cell in row.cells)
        except Exception as exc:
            raise ValueError("Could not read this DOCX file.") from exc
    elif suffix == ".txt":
        text = content.decode("utf-8", errors="replace")
    else:
        raise ValueError("Unsupported file type. Upload PDF, DOCX, or TXT.")
    if not text.strip():
        raise ValueError("No readable text was found. Scanned PDFs need OCR before analysis.")
    return text[:100_000]
